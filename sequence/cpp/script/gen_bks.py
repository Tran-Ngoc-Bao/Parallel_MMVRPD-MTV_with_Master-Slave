#!/usr/bin/env python3
"""
Generate bks/<n>/<instance>-bks.json for every instance in data/ from the
best solution recorded per problem in bks/capacity-1400_baseline.xlsx.

The xlsx only has routes + a self-reported "Cost [minute]" column, so this
script does NOT trust that number directly. Instead it uses the sequence
binary itself as the source of truth:

  1. `sequence run <problem> --dry-run --compact-output` to obtain a valid
     config skeleton (customer coordinates/demands/dronable + the default
     truck/drone physical parameters) for that instance, without running
     any search.
  2. `sequence evaluate <solution.json> <config.json>` with the xlsx
     truck/drone routes plugged into solution.json, so the solver itself
     computes working_time/feasibility/violations.

If bks/<n>/<instance>-bks.json already exists and its working_time is
lower (better) than the xlsx-derived candidate, the existing file is kept
untouched -- it may have come from an actual solver run found after the
xlsx baseline was recorded. Use --force to overwrite regardless.

Usage:
    python3 gen_bks.py                  # fill in / update all 128 instances
    python3 gen_bks.py --customers 6,10 # only these customer counts
    python3 gen_bks.py --dry-run        # print what would happen, write nothing
"""
import argparse
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

SCRIPT_DIR = Path(__file__).resolve().parent
CPP_DIR = SCRIPT_DIR.parent
PROJECT_ROOT = CPP_DIR.parent.parent
BUILD_BIN = CPP_DIR / "build" / "sequence"
DATA_DIR = PROJECT_ROOT / "data"
BKS_DIR = PROJECT_ROOT / "bks"
XLSX_PATH = BKS_DIR / "capacity-1400_baseline.xlsx"

# xlsx "Cost [minute]" vs working_time(seconds)/60 should match closely;
# anything further off than this is treated as a parsing/logic error.
COST_CHECK_REL_TOL = 1e-3


def load_best_rows(xlsx_path: Path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["capacity-1400"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    idx = {c: i for i, c in enumerate(header)}

    best = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        problem = row[idx["Problem"]]
        if not problem:
            continue
        cost = row[idx["Cost [minute]"]]
        if problem not in best or cost < best[problem]["cost"]:
            best[problem] = {
                "cost": cost,
                "truck_paths": json.loads(row[idx["Truck paths"]]),
                "drone_paths": json.loads(row[idx["Drone paths"]]),
            }
    return best


def run_cmd(cmd):
    # The binary loads problems/config_parameter/*.json relative to CWD.
    r = subprocess.run(cmd, capture_output=True, text=True, cwd=str(CPP_DIR))
    if r.returncode != 0:
        raise RuntimeError(
            f"command failed: {' '.join(cmd)}\nstdout={r.stdout}\nstderr={r.stderr}")
    return r


def build_skeleton_config(problem: str, tmpdir: Path):
    data_file = DATA_DIR / f"{problem}.txt"
    if not data_file.is_file():
        raise FileNotFoundError(data_file)
    skel_dir = tmpdir / "skel"
    skel_dir.mkdir(exist_ok=True)
    run_cmd([str(BUILD_BIN), "run", str(data_file), "--dry-run",
              "--outputs", str(skel_dir), "--run-id", "sk", "--compact-output"])
    skel_file = skel_dir / f"{problem}-sk.json"
    with skel_file.open() as f:
        return json.load(f)["config"]


def evaluate(problem: str, config: dict, truck_paths, drone_paths, tmpdir: Path):
    eval_dir = tmpdir / "eval"
    eval_dir.mkdir(exist_ok=True)
    config_path = tmpdir / f"{problem}-config.json"
    solution_path = tmpdir / f"{problem}-solution.json"

    config = dict(config)
    config["outputs"] = str(eval_dir)
    with config_path.open("w") as f:
        json.dump(config, f)
    with solution_path.open("w") as f:
        json.dump({"truck_routes": truck_paths, "drone_routes": drone_paths}, f)

    run_cmd([str(BUILD_BIN), "evaluate", str(solution_path), str(config_path)])

    candidates = [p for p in eval_dir.glob(f"{problem}-*.json")
                  if not p.name.endswith("-solution.json")
                  and not p.name.endswith("-config.json")]
    if len(candidates) != 1:
        raise RuntimeError(f"expected exactly 1 output file for {problem}, got {candidates}")
    with candidates[0].open() as f:
        return json.load(f)


def generate_one(problem: str, xlsx_entry: dict, tmpdir: Path):
    config = build_skeleton_config(problem, tmpdir)
    result = evaluate(problem, config, xlsx_entry["truck_paths"],
                       xlsx_entry["drone_paths"], tmpdir)

    working_time = result["solution"]["working_time"]
    feasible = result["solution"]["feasible"]

    expected_wt = xlsx_entry["cost"] * 60.0
    if expected_wt and abs(working_time - expected_wt) > COST_CHECK_REL_TOL * expected_wt:
        raise RuntimeError(
            f"{problem}: solver working_time={working_time} disagrees with "
            f"xlsx Cost*60={expected_wt} by more than {COST_CHECK_REL_TOL:.1%}")

    return result, working_time, feasible


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                  formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--customers", default=None,
                     help="Comma-separated customer counts to (re)generate "
                          "(default: all n found in data/)")
    ap.add_argument("--force", action="store_true",
                     help="Overwrite existing bks files even if their "
                          "working_time is already better")
    ap.add_argument("--dry-run", action="store_true",
                     help="Print what would happen, write nothing")
    args = ap.parse_args()

    if not BUILD_BIN.is_file():
        sys.exit(f"error: {BUILD_BIN} not found -- build the project first "
                  f"(sequence/cpp/script/build.sh)")

    print(f"Loading baseline routes from {XLSX_PATH} ...")
    best_rows = load_best_rows(XLSX_PATH)
    print(f"  {len(best_rows)} distinct problems in xlsx")

    all_problems = sorted(p.stem for p in DATA_DIR.glob("*.txt"))
    if args.customers:
        wanted_n = {c.strip() for c in args.customers.split(",") if c.strip()}
        all_problems = [p for p in all_problems if p.split(".")[0] in wanted_n]

    missing_in_xlsx = [p for p in all_problems if p not in best_rows]
    if missing_in_xlsx:
        print(f"WARNING: {len(missing_in_xlsx)} data instance(s) have no xlsx "
              f"baseline row, skipping: {missing_in_xlsx}")
        all_problems = [p for p in all_problems if p in best_rows]

    kept, written, failed = [], [], []

    with tempfile.TemporaryDirectory(prefix="gen_bks_") as tmp:
        tmpdir = Path(tmp)
        for problem in all_problems:
            n = problem.split(".")[0]
            bks_path = BKS_DIR / n / f"{problem}-bks.json"

            existing_wt = None
            if bks_path.is_file():
                with bks_path.open() as f:
                    existing = json.load(f)
                if existing["solution"]["feasible"]:
                    existing_wt = existing["solution"]["working_time"]

            try:
                result, working_time, feasible = generate_one(
                    problem, best_rows[problem], tmpdir)
            except Exception as e:
                print(f"FAIL  {problem}: {e}")
                failed.append(problem)
                continue

            if not feasible:
                print(f"FAIL  {problem}: xlsx-derived solution is infeasible per solver")
                failed.append(problem)
                continue

            if existing_wt is not None and not args.force and existing_wt <= working_time:
                print(f"keep  {problem}: existing bks={existing_wt:.4f} "
                      f"<= xlsx-derived={working_time:.4f}")
                kept.append(problem)
                continue

            reason = "new" if existing_wt is None else f"better than existing {existing_wt:.4f}"
            print(f"WRITE {problem}: working_time={working_time:.4f} ({reason})")
            written.append(problem)

            if not args.dry_run:
                bks_path.parent.mkdir(parents=True, exist_ok=True)
                with bks_path.open("w") as f:
                    json.dump(result, f)

    print(f"\nDone: {len(written)} written, {len(kept)} kept as-is, "
          f"{len(failed)} failed, out of {len(all_problems)} instances.")
    if failed:
        print(f"Failed: {failed}")
        sys.exit(1)


if __name__ == "__main__":
    main()
